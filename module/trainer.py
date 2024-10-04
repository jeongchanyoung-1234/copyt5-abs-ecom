# import re, os, operator
# from queue import PriorityQueue
from string import ascii_uppercase
from datetime import datetime

#import mecab
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from tqdm import tqdm
from nltk.translate.bleu_score import sentence_bleu, SmoothingFunction

import torch
import torch.nn as nn
import torch.nn.functional as F

from module.utils import eval_bleu, bleu_score


class Trainer:
    def __init__(self, 
                model, 
                tokenizer,
                optim, 
                scheduler,
                train_loader,
                test_loader, 
                config,
                device):
        self.model = model
        self.tokenizer = tokenizer
        self.optim = optim
        self.scheduler = scheduler
        self.lr = config.lr
        self.train_loader = train_loader
        self.test_loader = test_loader
        self.train_data_path = config.train_data_path
        self.valid_data_path = config.test_data_path
        self.device = device
        self.config = config

        self.epochs = config.epochs
        self.num_beams = config.num_beams
        self.lr_decay_factor = config.lr_decay_factor
        self.max_target_length = config.max_target_length
        self.warmup_stage = config.warmup_stage
        self.early_stopping_round = config.early_stopping_round
        self.early_stopping_cnt = config.early_stopping_round
        self.early_stopping_criterion = config.early_stopping_criterion
        self.save_path = config.save_path
        self.verbose = config.verbose

        self.best_score = 0.
        try:
            self.vocab_size = model.generator.config.vocab_size
            self.hidden_dim = model.generator.config.d_model
        except:
            self.vocab_size = model.generator.encoder.config.vocab_size
            self.hidden_dim = model.generator.encoder.config.hidden_size

        self.pad_token_id = model.generator.config.pad_token_id
        self.eos_token_id = model.generator.config.eos_token_id
        self.epsilon = 1e-6

        #self.mecab = mecab.MeCab()
        self.use_copy = config.use_copy
        self.gen_loss_lambda = config.gen_loss_lambda
        self.model.on_device(self.device)
        self.coverage = config.coverage

        self.max_grad_norm = config.max_grad_norm
        self.print_example = config.print_example
        self.strategy = config.tokenizing_strategy

        self.print_test_result = config.print_test_result
        self.save_result_to_xlsx = config.save_result_to_xlsx

        self.scaler = None
        if config.fp16:
            self.scaler = torch.cuda.amp.GradScaler()

    def split_tokens(self, sentence):
        if self.strategy == 'space':
            return sentence.split(' ')
        #elif strategy == 'mecab':
        #    return self.mecab.morphs(sentence)
        else:
            return [sentence]
    
    def init_workbook(self):
        wb = Workbook()
        wb.create_sheet('Model Info')
        result_sheet = wb.active
        result_sheet.title = 'Result'
        
        result_sheet['A1'] = 'Input ID'
        result_sheet['B1'] = 'Reference'
        result_sheet['C1'] = 'Hypothesis'
        result_sheet['D1'] = 'Is Matched'
        result_sheet['E1'] = 'Statistics'

        colfill = PatternFill(start_color='cfff7d', end_color='cfff7d', fill_type='solid')
        for char in ascii_uppercase[:5]: # A~E
            result_sheet['{}1'.format(char)].fill = colfill

        model_sheet = wb['Model Info']
        config_dic = vars(self.config)
        for i, (k, v) in enumerate(config_dic.items()):
            model_sheet['A{}'.format(i + 1)] = k
            model_sheet['B{}'.format(i + 1)] = v
        return wb, result_sheet
    
    def _shift_right(self, input_ids):
        decoder_start_token_id = self.model.generator.config.decoder_start_token_id
        pad_token_id = self.pad_token_id

        assert (
            decoder_start_token_id is not None
        ), "self.model.config.decoder_start_token_id has to be defined. In T5 it is usually set to the pad_token_id. See T5 docs for more information"

        # shift inputs to the right
        shifted_input_ids = input_ids.new_zeros(input_ids.shape)
        shifted_input_ids[..., 1:] = input_ids[..., :-1].clone()
        shifted_input_ids[..., 0] = decoder_start_token_id

        assert pad_token_id is not None, "self.model.config.pad_token_id has to be defined."
        # replace possible -100 values in labels by `pad_token_id`
        shifted_input_ids.masked_fill_(shifted_input_ids == -100, pad_token_id)

        assert torch.all(shifted_input_ids >= 0).item(), "Verify that `labels` has only positive values and -100"

        return shifted_input_ids

    def gradient_clipping(self):
        if self.use_copy:
            nn.utils.clip_grad_norm_(self.model.generator.parameters(), self.max_grad_norm)
            nn.utils.clip_grad_norm_(self.model.attention.parameters(), self.max_grad_norm)
            nn.utils.clip_grad_norm_(self.model.p_gen.parameters(), self.max_grad_norm)
        else:
            nn.utils.clip_grad_norm_(self.model.generator.parameters(), self.max_grad_norm)
    
    def save(self, epoch):
        if self.use_copy:
            checkpoint = {
            "generator": self.model.generator.state_dict(),
            "attention": self.model.attention.state_dict(),
            "p_gen": self.model.p_gen.state_dict(),
            "optimizer": self.optim.state_dict(),
            "epoch": epoch,
            "score": 0,
            "config": self.config
            }
        else:
            checkpoint = {
            "generator": self.model.generator.state_dict(),
            "optimizer": self.optim.state_dict(),
            "epoch": epoch,
            "score": 0,
            "config": self.config
            }
        torch.save(checkpoint, self.save_path)
        save_fn = self.save_path
        if self.early_stopping_round <= 0:
            save_fn = save_fn.split('.')
            save_fn[1] = save_fn[1] + '_ep{}'.format(epoch)
            save_fn = '.'.join(save_fn)
        print('saved to {}'.format(save_fn))

    def generate(self, input_ids):
        batch_size = input_ids.size(0)
        input_ids = input_ids.to(self.device, dtype=torch.long)
        mask = (input_ids == self.pad_token_id).long()
        i, coverage, hypothesis = 0, None, []

        is_decoding = input_ids.new_ones(batch_size, 1).bool()
        decoder_input_ids = input_ids.new_zeros(batch_size, 1) + self.model.generator.config.decoder_start_token_id

        while is_decoding.sum() > 0 and len(hypothesis) < self.max_target_length:
            outputs = self.model.generator.generate(input_ids=input_ids,
                                                    attention_mask=(1 - mask),
                                                    decoder_input_ids=decoder_input_ids,
                                                    max_new_tokens=1,
                                                    return_dict_in_generate=True,
                                                    output_scores=True,
                                                    num_beams=self.num_beams,
                                                    output_hidden_states=True)

            if self.use_copy:
                scores = torch.stack(outputs.scores, dim=1)
                h_t_dec = outputs.decoder_hidden_states[0][-1][:, i, :]
                p_gen_history = outputs.decoder_hidden_states[0][0][:, i, :]
                h_enc = outputs.encoder_hidden_states[-1]
                vocab_dist = F.softmax(scores[:, 0, :], dim=-1)  # (B, VOCAB)

                if self.coverage:
                    if coverage is None:
                        coverage = torch.zeros(batch_size, h_enc.size(1)).to(self.device) if self.coverage else None
                    z_enc = coverage
                else:
                    z_enc = None
                c_t, attn_dist = self.model.attention(h_t_dec, h_enc, z_enc=z_enc, mask=mask)
                p_gen = self.model.p_gen(c_t, h_t_dec, p_gen_history)

                if self.coverage:
                    coverage = coverage + attn_dist

                vocab_dist = p_gen * vocab_dist
                attn_dist = (1 - p_gen) * attn_dist

                outputs = vocab_dist.scatter_add(1, input_ids, attn_dist)  # (B, VOCAB)
            else:
                scores = torch.stack(outputs.scores, dim=1)
                outputs = scores[:, -1, :]

            y = outputs.argmax(dim=-1).unsqueeze(1)
            y = y.masked_fill_(~is_decoding, self.pad_token_id)
            is_decoding = is_decoding * torch.ne(y, self.eos_token_id)
            hypothesis += [y]
            decoder_input_ids = torch.cat([decoder_input_ids, y.view(batch_size, 1).contiguous()], dim=1)
            i += 1
        hypothesis = torch.cat(hypothesis, dim=1)
        hypothesis = self.tokenizer.batch_decode(hypothesis, skip_special_tokens=True)
        hypothesis = list(map(str.strip, hypothesis))

        hypothesis = hypothesis[0]

        return hypothesis


    def train(self, global_epoch=0):
        self.model.run_train()
        for epoch in range(global_epoch, self.epochs + global_epoch):
            epoch += 1
            loop = tqdm(self.train_loader, leave=True, miniters=10) if self.verbose >= 1 else self.train_loader
            total_loss = 0.
            loss_cnt = 0
            for input_ids, labels in loop:
                self.optim.zero_grad()
            
                input_ids = input_ids.to(self.device, dtype=torch.long)
                mask = (input_ids == self.pad_token_id).long()
                labels = labels.to(self.device, dtype=torch.long)
                decoder_input_ids = self._shift_right(labels)
                batch_size = input_ids.size(0)

                outputs = self.model.generator(
                    input_ids=input_ids,
                    attention_mask=(1 - mask),
                    decoder_input_ids=decoder_input_ids,
                    labels=labels,
                    output_hidden_states=True)

                if self.use_copy:
                    gen_loss = torch.tensor([[0]]).float().to(self.device)
                    gen_loss_cnt = 0
                    total_cnt = 0

                    step_losses = []

                    h_enc = outputs.encoder_last_hidden_state
                    
                    coverage = torch.zeros(batch_size, h_enc.size(1)).to(self.device) if self.coverage else None
                    coverage_loss = []

                    h_dec = outputs.decoder_hidden_states[-1]
                    final_dists = []
                    for i in range(h_dec.size(1)):
                        h_t_dec = h_dec[:, i, :]
                        target = labels[:, i]
                        p_gen_history = outputs.decoder_hidden_states[0][:, i, :]
                        vocab_dist = F.softmax(outputs.logits[:, i, :], dim=-1)

                        if self.coverage:
                            z_enc = coverage
                        else:
                            z_enc = None
                        c_t, attn_dist = self.model.attention(h_t_dec, h_enc, z_enc=z_enc, mask=mask)
                        attn_dist = attn_dist * (1 - mask)

                        p_gen = self.model.p_gen(c_t, h_t_dec, p_gen_history) 
                        if self.coverage:
                            cov_loss = torch.minimum(attn_dist, coverage).sum(1)
                            coverage_loss.append(cov_loss)
                            coverage = coverage + attn_dist

                        for b in range(batch_size):
                            total_cnt += 1
                            if target[b].item() in input_ids[b].tolist():
                                gen_loss += p_gen[b]
                                gen_loss_cnt += 1

                        # if epoch % 2 == 0:
                        #     p_gen = 1
                        vocab_dist = p_gen * vocab_dist
                        attn_dist = (1 - p_gen) * attn_dist
 
                        final_dist = vocab_dist.scatter_add(1, input_ids, attn_dist)
                        final_dists.append(final_dist.unsqueeze(1))
                        gold_probs = torch.gather(final_dist, 1, target.unsqueeze(1)).squeeze()
                        step_loss = -torch.log(torch.clamp(gold_probs, min=self.epsilon))
                        step_losses.append(step_loss)

                    batch_avg_loss = torch.mean(torch.stack(step_losses, -1), -1)
                    loss = torch.mean(batch_avg_loss)
                    
                    gen_loss = gen_loss / gen_loss_cnt
                    gen_loss_lambda = gen_loss_cnt / total_cnt
                    # gen_loss = gen_loss_lambda * gen_loss[0][0]
                    gen_loss = self.gen_loss_lambda * gen_loss[0][0]
                    loss += gen_loss
                    if self.coverage:
                        cov_loss_cnt = len(coverage_loss)
                        coverage_loss = torch.sum(torch.stack(coverage_loss, 1), 1)
                        coverage_loss = torch.mean(coverage_loss / cov_loss_cnt)
                        loss += coverage_loss

                else:
                    loss = F.cross_entropy(outputs.logits.view(-1, self.vocab_size),
                                           labels.contiguous().view(-1), ignore_index=self.pad_token_id)
                    
                if self.scaler is not None:
                    self.scaler.scale(loss).backward()
                    self.gradient_clipping()
                    self.scaler.step(self.optim)
                    self.scaler.update()
                else:
                    loss.backward()
                    self.gradient_clipping()
                    self.optim.step()
                    
                if self.scheduler is not None:
                    self.scheduler.step()

                loss = loss.detach()

                if self.verbose > 0:
                    loop.set_description('Epoch ({}) loss {:.4f}'.format(epoch, loss))
                
                total_loss += loss
                loss_cnt += 1

            torch.cuda.empty_cache()

            total_loss = total_loss / loss_cnt
            print('[TRAINING] loss: {:.4f}'.format(total_loss))
            
            if epoch > self.warmup_stage:
                if self.early_stopping_round > 0:
                    self.validate(epoch)
                else:
                    self.save(epoch)
            else:
                print('Warmup stage is set to {}. skip validation.'.format(self.warmup_stage))

    def validate(self, epoch, is_test=False):
        if self.save_result_to_xlsx:
            xlsx_offset, cmplt_match_cnt = 2, 0
            wb, sheet = self.init_workbook()
            

        self.model.run_eval()
        with torch.no_grad():
            total_bleu = 0.
            len_hyps, len_refs, cnt = 0, 0, 0
            t = tqdm(self.test_loader, miniters=1, leave=True) if self.verbose >= 1 else self.test_loader
            predictions = []
            self.optim.zero_grad()

            for input_ids, labels in t:
                batch_size = input_ids.size(0)
                input_ids = input_ids.to(self.device, dtype=torch.long)
                mask = (input_ids == self.pad_token_id).long()
                labels = labels.to(self.device, dtype=torch.long)
                i, coverage, hypothesis = 0, None, []

                if self.num_beams <= 1:
                    is_decoding = input_ids.new_ones(batch_size, 1).bool()
                    decoder_input_ids = input_ids.new_zeros(batch_size, 1) + self.model.generator.config.decoder_start_token_id
                    
                    while is_decoding.sum() > 0 and len(hypothesis) < self.max_target_length:
                        outputs = self.model.generator.generate(input_ids=input_ids,
                                                                attention_mask=(1 - mask),
                                                                decoder_input_ids=decoder_input_ids,
                                                                max_new_tokens=1,
                                                                return_dict_in_generate=True,
                                                                output_scores=True,
                                                                num_beams=self.num_beams,
                                                                output_hidden_states=True)

                        if self.use_copy:
                            scores = torch.stack(outputs.scores, dim=1)
                            h_t_dec = outputs.decoder_hidden_states[0][-1][:, i, :]
                            p_gen_history = outputs.decoder_hidden_states[0][0][:, i, :]
                            h_enc = outputs.encoder_hidden_states[-1]
                            vocab_dist = F.softmax(scores[:, 0, :], dim=-1) #(B, VOCAB)

                            if self.coverage:
                                if coverage is None:
                                    coverage = torch.zeros(batch_size, h_enc.size(1)).to(self.device) if self.coverage else None
                                z_enc = coverage
                            else:
                                z_enc = None
                            c_t, attn_dist = self.model.attention(h_t_dec, h_enc, z_enc=z_enc, mask=mask)
                            p_gen = self.model.p_gen(c_t, h_t_dec, p_gen_history)

                            if self.coverage:
                                coverage = coverage + attn_dist

                            vocab_dist = p_gen * vocab_dist
                            attn_dist = (1 - p_gen) * attn_dist

                            outputs = vocab_dist.scatter_add(1, input_ids, attn_dist) #(B, VOCAB)
                        else:
                            scores = torch.stack(outputs.scores, dim=1)
                            outputs = scores[:, -1, :]

                        y = outputs.argmax(dim=-1).unsqueeze(1)
                        y = y.masked_fill_(~is_decoding, self.pad_token_id)
                        is_decoding = is_decoding * torch.ne(y, self.eos_token_id)
                        hypothesis += [y]
                        decoder_input_ids = torch.cat([decoder_input_ids, y.view(batch_size, 1).contiguous()], dim=1)
                        i += 1
                    hypothesis = torch.cat(hypothesis, dim=1)
                else:
                    done_cnt = 0
                    cumulative_probs = torch.zeros(batch_size, self.num_beams).to(self.device)
                    ys = torch.zeros(batch_size, self.num_beams, 1).long().to(self.device)
                    decoder_input_ids = input_ids.new_zeros(batch_size * self.num_beams, 1) + self.model.generator.config.decoder_start_token_id

                    input_ids_ = input_ids.unsqueeze(0).repeat(self.num_beams, 1, 1).transpose(0, 1).flatten(end_dim=-2)
                    mask = (input_ids_ == self.pad_token_id).long()
                    #|input_ids|, |labels| = (b * n_beams, len)
                    while done_cnt < batch_size * self.num_beams and i < self.max_target_length:
                        outputs = self.model.generator.generate(input_ids=input_ids_,
                                                                attention_mask=(1 - mask),
                                                                decoder_input_ids=decoder_input_ids,
                                                                max_new_tokens=1,
                                                                return_dict_in_generate=True, 
                                                                output_scores=True,
                                                                output_hidden_states=True)

                        if self.use_copy:
                            scores = torch.stack(outputs.scores, dim=1)
                            h_t_dec = outputs.decoder_hidden_states[0][-1][:, i, :]
                            p_gen_history = outputs.decoder_hidden_states[0][0][:, i, :]
                            h_enc = outputs.encoder_hidden_states[-1]
                            
                            vocab_dist = F.softmax(scores[:, 0, :], dim=-1) #(B, VOCAB)
                            
                            if self.coverage:
                                if coverage is None:
                                    coverage = torch.zeros(batch_size * self.num_beams, h_enc.size(1)).to(self.device) if self.coverage else None
                                z_enc = coverage
                            else:
                                z_enc = None

                            c_t, attn_dist = self.model.attention(h_t_dec, h_enc, z_enc=z_enc, mask=mask)
                            # attn_dist = attn_dist * (1 - mask)
                            p_gen = self.model.p_gen(c_t, h_t_dec, p_gen_history)

                            if self.coverage:
                                coverage = coverage + attn_dist

                            vocab_dist = p_gen * vocab_dist
                            attn_dist = (1 - p_gen) * attn_dist

                            outputs = vocab_dist.scatter_add(1, input_ids, attn_dist) #(b * n_beams, VOCAB)
                        else:
                            outputs = outputs.logits[:, -1, :]
                        cumulative_probs = cumulative_probs.view(-1, 1) + torch.log(outputs) #(b * n_beams, vocab)
                        cumulative_probs = cumulative_probs.view(batch_size, -1) #(b, n_beams * vocab)
                        cumulative_probs, indices = torch.topk(cumulative_probs, self.num_beams, dim=-1) #(b, n_beams)

                        prev_beam_indices = indices // self.vocab_size
                        next_beam_indices = indices - self.vocab_size * prev_beam_indices
                        beam_mask = (ys[:, :, i] == self.eos_token_id) #b, n_beams
                        next_beam_indices[beam_mask] = self.eos_token_id
                        ys[:, :, i] = torch.gather(ys[:, :, i], 1, prev_beam_indices)

                        ys = torch.concat([ys, next_beam_indices.unsqueeze(-1)], dim=-1) #(b, n_beams, len)
                        decoder_input_ids = ys.flatten(end_dim=-2)
                        done_cnt = torch.sum(next_beam_indices == self.eos_token_id)
                        i += 1

                    indices = cumulative_probs.view(batch_size, self.num_beams).topk(1, dim=-1).indices.squeeze(-1) #(b)
                    for b in range(batch_size):
                        hypothesis.append(ys[b, indices[b]])
                    hypothesis = torch.stack(hypothesis, dim=0)

                bleu = 0
                hypothesis = self.tokenizer.batch_decode(hypothesis, skip_special_tokens=True)
                hypothesis = list(map(str.strip, hypothesis))
                labels = self.tokenizer.batch_decode(labels, skip_special_tokens=True)
                labels = list(map(str.strip, labels))
                for b in range(batch_size):
                    hypothesis_decoded = hypothesis[b]
                    reference = labels[b]
                    if self.print_example or self.save_result_to_xlsx:
                        input_decoded = self.tokenizer.decode(input_ids[b], skip_special_tokens=True)
                        if self.print_example:
                            print('input_ids\n', input_decoded)
                            print('reference\n', reference)
                            print('hypothesis_decoded\n', hypothesis_decoded)
                            print('============================================')
                        if self.save_result_to_xlsx:
                            sheet['A{}'.format(xlsx_offset)] = input_decoded
                            sheet['B{}'.format(xlsx_offset)] = reference
                            sheet['C{}'.format(xlsx_offset)] = hypothesis_decoded
                            if reference.strip() == hypothesis_decoded.strip():
                                sheet['D{}'.format(xlsx_offset)] = 'True'
                                cmplt_match_cnt += 1
                            else:
                                sheet['D{}'.format(xlsx_offset)] = 'False'
                            xlsx_offset += 1

                    len_hypothesis = len(hypothesis_decoded)
                    hypothesis_tok = self.split_tokens(hypothesis_decoded)
                    len_reference = len(reference)
                    if self.print_test_result:
                        predictions.append(hypothesis_decoded)
                    reference_tok = [self.split_tokens(reference)]
                    
                    bleu += sentence_bleu(reference_tok, hypothesis_tok, weights=(.25, .25, .25, .25), smoothing_function=SmoothingFunction().method1)
                    # bleu += round(corpus_bleu(hypothesis_tok, reference_tok).score, 4)
                    # bleu += metric(hypothesis_tok, reference_tok)

                bleu = bleu / batch_size

                total_bleu += bleu
                len_hyps += len_hypothesis
                len_refs += len_reference
                cnt += 1

                if self.verbose > 0:
                    t.set_description('Epoch ({}) bleu {:.4f}'.format(epoch, total_bleu / cnt))
            
            torch.cuda.empty_cache()

            total_bleu = total_bleu / cnt
            mean_len_hyps = len_hyps / cnt
            mean_len_refs = len_refs / cnt
            
            print('[VALIDATION] \nbleu: {:.4f}'.format(total_bleu))
            print('Length: hyp={:.4f} ref={:.4f}'.format(mean_len_hyps, mean_len_refs))
            
            if self.save_result_to_xlsx:
                test_len = len(self.test_loader) * batch_size
                sheet['E2'] = 'Exactly Matched: {}'.format(cmplt_match_cnt)
                sheet['E3'] = 'Total: {}'.format(test_len)
                sheet['E4'] = 'Accuracy: {:.2f}'.format(cmplt_match_cnt / test_len)
                result_save_path = 'C:\\Users\\builton\\Desktop\\abstract_model\\result\\{}.{}.xlsx'.format(
                    self.save_path.split('/')[-1].split('.')[0],
                    datetime.today().strftime('%Y-%m-%d_%H-%M-%S'))
                wb.save(result_save_path)
                wb.close()
                print('Result file successfully saved to {}'.format(result_save_path))

            if is_test:
                exit()
                
            criterion_map = {'bleu': total_bleu}
            score = criterion_map[self.early_stopping_criterion]
            if score > self.best_score:
                self.early_stopping_cnt = self.early_stopping_round
                print(' best score updated, saved at {}'.format(self.save_path))
                self.best_score = score
                checkpoint = {
                "generator": self.model.generator.state_dict(),
                "attention": self.model.attention.state_dict(),
                "p_gen": self.model.p_gen.state_dict(),
                "optimizer": self.optim.state_dict(),
                "epoch": epoch,
                "score": score,
                "config": self.config
                } if self.use_copy else {
                "generator": self.model.generator.state_dict(),
                "optimizer": self.optim.state_dict(),
                "epoch": epoch,
                "score": score,
                "config": self.config
                }
                torch.save(checkpoint, self.save_path)
            else:
                self.early_stopping_cnt -= 1
                if self.lr_decay_factor > 0:
                    self.lr /= self.lr_decay_factor
                    for param in self.optim.param_groups:
                        param["lr"] = self.lr
                    print('Learning rate decayed to {}'.format(self.lr))
                if self.early_stopping_cnt == 0:
                    print(' Best score not updated during {} epochs, training stopped'.format(self.early_stopping_round))
                    exit()